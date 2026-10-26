import argparse,csv,json,hashlib,sqlite3,html,math
from pathlib import Path
FIELDS=['plant_id','plant_name','state','generator_id','capacity_mw','energy_source']
ALIASES=dict(zip(['Plant Code','Plant Name','State','Generator ID','Nameplate Capacity (MW)','Energy Source 1'],FIELDS))
def analyze(rows):
 clean=[];issues=[];seen={}
 for n,raw in rows:
  mapped={ALIASES.get(k.strip(),k.strip()):v for k,v in raw.items()}
  if not set(FIELDS)<=mapped.keys(): raise ValueError('Missing required fields: '+str(set(FIELDS)-mapped.keys()))
  r={k:str(mapped[k]).strip() if mapped[k] is not None else '' for k in FIELDS};r['source_row']=n
  def flag(field,rule,value):issues.append(dict(source_row=n,field=field,rule=rule,value=value))
  for k in FIELDS:
   if not r[k]:flag(k,'missing',r[k])
  key=(r['plant_id'],r['generator_id'])
  if all(key):
   if key in seen:flag('generator_key','duplicate',f'first row {seen[key]}')
   else:seen[key]=n
  if r['state'] not in 'AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY AS GU MP PR VI'.split():flag('state','invalid_state',r['state'])
  try:
   v=float(r['capacity_mw'])
   if not math.isfinite(v) or v<=0:flag('capacity_mw','invalid_capacity',r['capacity_mw'])
   elif v>2000:flag('capacity_mw','review_large_generator',r['capacity_mw'])
  except ValueError:flag('capacity_mw','not_numeric',r['capacity_mw'])
  clean.append(r)
 return clean,issues

def load(p,sheet):
 if p.suffix.lower()=='.csv':
  with p.open(encoding='utf-8-sig',newline='') as f:return list(enumerate(csv.DictReader(f),2))
 import openpyxl
 w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 if sheet not in w.sheetnames:raise ValueError(f'Choose a sheet from {w.sheetnames}')
 it=w[sheet].iter_rows(values_only=True)
 for n,row in enumerate(it,1):
  headers=[str(v).strip() if v is not None else '' for v in row]
  if 'Plant Code' in headers and 'Generator ID' in headers:return [(i,dict(zip(headers,r))) for i,r in enumerate(it,n+1) if any(v is not None for v in r) and not (str(r[0]).startswith('NOTE:') and all(v is None for v in r[1:]))]
  if n==20:break
 raise ValueError('Generator header not found in first 20 rows')
def main():
 p=argparse.ArgumentParser();p.add_argument('input',type=Path);p.add_argument('--source-url',required=True);p.add_argument('--sheet',default='Operable');p.add_argument('--demo',action='store_true');p.add_argument('--out',type=Path,default=Path('output'));a=p.parse_args()
 clean,issues=analyze(load(a.input,a.sheet));a.out.mkdir(parents=True,exist_ok=True)
 for name,data,fields in [('normalized',clean,FIELDS+['source_row']),('issues',issues,['source_row','field','rule','value'])]:
  with (a.out/(name+'.csv')).open('w',newline='') as f:
   w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(data)
 count=len({i['source_row'] for i in issues})
 m=dict(mode='SYNTHETIC DEMO' if a.demo else 'USER-SUPPLIED SOURCE',source_url=a.source_url,source_file=a.input.name,sheet=a.sheet if a.input.suffix!='.csv' else None,sha256=hashlib.sha256(a.input.read_bytes()).hexdigest(),rows=len(clean),flagged_rows=count,issue_count=len(issues),flagged_fraction=count/len(clean) if clean else None)
 (a.out/'manifest.json').write_text(json.dumps(m,indent=2))
 with sqlite3.connect(a.out/'gridcheck.sqlite') as db:
  db.execute('DROP TABLE IF EXISTS generators');db.execute('CREATE TABLE generators ('+', '.join(k+' TEXT' for k in FIELDS)+',source_row INTEGER PRIMARY KEY)');db.executemany('INSERT INTO generators VALUES (?,?,?,?,?,?,?)',[[r[k] for k in FIELDS]+[r['source_row']] for r in clean])
 e=lambda x:html.escape(str(x))
 report='<html><meta charset="utf-8"><title>GRIDCHECK</title><style>body{font:16px system-ui;background:#eff5f2;color:#143e35;max-width:1000px;margin:50px auto;padding:24px}h1{font-size:48px}table{background:white;width:100%;border-collapse:collapse}td,th{padding:12px;text-align:left;border-bottom:1px solid #ddd}pre{white-space:pre-wrap;overflow-wrap:anywhere}</style><p>ENERGY DATA QUALITY</p><h1>GRIDCHECK</h1>'
 report+='<h2>'+e(m['mode'])+'</h2><p>'+f'{len(clean)} records · {count} flagged records · {len(issues)} issues'+'</p><p>Flags require review; they are not confirmed errors. All rows are retained. Capacity is MW, not energy generated.</p><table><tr><th>Source row</th><th>Field</th><th>Rule</th><th>Evidence</th></tr>'
 report+=''.join('<tr>'+''.join('<td>'+e(r[k])+'</td>' for k in ['source_row','field','rule','value'])+'</tr>' for r in issues)
 report+='</table><h2>Provenance</h2><pre>'+e(json.dumps(m,indent=2))+'</pre><p>Large generator threshold: 2,000 MW (heuristic). No runtime LLM inference or investment recommendations. Generated with Codex assistance.</p></html>'
 (a.out/'report.html').write_text(report);print(json.dumps(m,indent=2))
if __name__=='__main__':main()
